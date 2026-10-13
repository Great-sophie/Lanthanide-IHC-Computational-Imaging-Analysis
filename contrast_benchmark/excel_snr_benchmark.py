#!/usr/bin/env python3
"""Audit existing Chapter 2 Excel SNR tables and create descriptive Eu-vs-Cy3 figure.

Reads the reported Excel columns; does NOT quantify rendered RGB TIFF images.
ROIs/rows must NOT automatically be interpreted as independent biological replicates.
"""
from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path
from statistics import mean, median, stdev

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

# Spreadsheet columns are ZERO-based (Excel column A = 0).
# Confirmed from the user's preview of the four historical spreadsheets.
SOURCES = [
    {'condition': 'BT474 / FFEP', 'mode': 'Cy3',
     'relative_path': '20220428-ihc-cy3/474-IHC-FFEP-SNR-cy3.xlsx',
     'columns': (1, 2, 3, 4)},
    {'condition': 'BT474 / FFEP', 'mode': 'Eu',
     'relative_path': '20220428-ihc-eu/474-TIHC-FFEP-SNR.xlsx',
     'columns': (0, 1, 2, 3)},
    {'condition': 'CK18 / FS', 'mode': 'Cy3',
     'relative_path': '20220428-ihc-cy3/IHC-FS-SNR.xlsx',
     'columns': (1, 2, 3, 4)},
    {'condition': 'CK18 / FS', 'mode': 'Eu',
     'relative_path': '20220428-ihc-eu/CK18-TIHC-FS-SNR.xlsx',
     'columns': (0, 1, 2, 3)},
]


def is_number(v):
    return isinstance(v, (float, int)) and not isinstance(v, bool) and math.isfinite(float(v))


def extract_rows(sheet_rows, source):
    """Read the four named columns; skip header and blank rows (never interpolate)."""
    idx_i, idx_s, idx_n, idx_r = source['columns']
    rows = []
    for row_number, values in enumerate(sheet_rows, start=1):
        vals = list(values)
        if len(vals) <= max(idx_i, idx_s, idx_n, idx_r):
            continue
        identifier, sig, noise, rep = (vals[i] for i in (idx_i, idx_s, idx_n, idx_r))
        if not all(is_number(x) for x in (identifier, sig, noise, rep)):
            continue
        if noise <= 0:
            raise ValueError(f"{source['relative_path']} row {row_number}: noise <= 0")
        if not float(identifier).is_integer():
            continue
        computed = float(sig) / float(noise)
        # With one-decimal reported SNR, 0.05 is the usual rounding tolerance.
        # For two decimals, the tolerance is 0.005. We flag but never overwrite.
        reported_text = str(rep)
        precision = len(reported_text.split('.', 1)[1]) if '.' in reported_text else 0
        tolerance = 0.5 * 10**(-precision) + 1e-9
        rows.append({
            'condition': source['condition'], 'modality': source['mode'],
            'workbook': source['relative_path'], 'excel_row': row_number,
            'measurement_id': int(identifier), 'signal': float(sig),
            'noise': float(noise), 'reported_snr': float(rep),
            'recomputed_signal_over_noise': computed,
            'absolute_difference': abs(float(rep) - computed),
            'rounding_tolerance': tolerance,
            'rounding_consistent': 'yes' if abs(float(rep)-computed) <= tolerance else 'no',
        })
    return rows


def write_csv(path, records, fields):
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
        w.writeheader()
        w.writerows(records)


def summarize(rows):
    result=[]
    for source in SOURCES:
        subset=[r for r in rows if r['condition'] == source['condition'] and r['modality'] == source['mode']]
        if not subset:
            raise ValueError(f"No numeric rows for {source['relative_path']}")
        vals=[r['reported_snr'] for r in subset]
        recalculated=[r['recomputed_signal_over_noise'] for r in subset]
        backgrounds=[r['noise'] for r in subset]
        result.append({
            'condition':source['condition'],'modality':source['mode'],
            'n_measurements':len(subset),'reported_snr_mean':mean(vals),
            'reported_snr_sd':stdev(vals) if len(vals)>1 else float('nan'),
            'reported_snr_median':median(vals),
            'recomputed_snr_mean':mean(recalculated),
            'noise_unique_values':len(set(backgrounds)),
            'rounding_inconsistencies':sum(r['rounding_consistent']=='no' for r in subset),
            'workbook':source['relative_path']
        })
    return result


def plot(rows, summary, out):
    conditions=list(dict.fromkeys(s['condition'] for s in summary))
    groups=[(c,m) for c in conditions for m in ('Cy3','Eu')]
    colors={'Cy3':'#536779', 'Eu':'#D47735'}
    fig, ax=plt.subplots(figsize=(9,5.3))
    for j,(condition,modality) in enumerate(groups):
        vals=[r['reported_snr'] for r in rows if r['condition']==condition and r['modality']==modality]
        x=np.full(len(vals), j, dtype=float)+np.linspace(-.085,.085,len(vals))
        ax.scatter(x, vals, s=28, alpha=.72, color=colors[modality], zorder=3)
        s=next(s for s in summary if s['condition']==condition and s['modality']==modality)
        ax.errorbar(j, s['reported_snr_mean'], yerr=s['reported_snr_sd'], fmt='D',
                    markersize=7, color='black', capsize=5, linewidth=1.8, zorder=5)
        ax.text(j, max(vals)+.24, f"{s['reported_snr_mean']:.2f}", ha='center', fontsize=10)
    for j,c in enumerate(conditions):
        cy=next(s['reported_snr_mean'] for s in summary if s['condition']==c and s['modality']=='Cy3')
        eu=next(s['reported_snr_mean'] for s in summary if s['condition']==c and s['modality']=='Eu')
        ax.text(2*j+.5, -0.13, f"Eu/Cy3 = {eu/cy:.2f}×", ha='center', va='top',
                transform=ax.get_xaxis_transform(), fontsize=10, fontweight='bold')
        if j+1<len(conditions):ax.axvline(2*j+1.5, color='0.83', linewidth=1)
    ax.set_xticks(range(len(groups)),[f'{m}\n{c}' for c,m in groups])
    ax.set_ylim(bottom=0,top=max(r['reported_snr'] for r in rows)+1.25)
    ax.set_ylabel('Reported Signal / Noise (historical Excel SNR)')
    ax.set_title('Chapter 2 | Historical Eu vs Cy3 contrast measurements', fontsize=14)
    ax.grid(axis='y',alpha=.2)
    fig.text(.5,.015,'Individual points = recorded measurements, not verified independent biological replicates.\n'
             'Different acquisition scales; interpret ratios only if measurement protocols are comparable.',
             ha='center', fontsize=8.6, color='#555555')
    fig.subplots_adjust(bottom=.28,top=.87,left=.11,right=.96)
    fig.savefig(out/'eu_vs_cy3_historical_snr.png',dpi=200)
    plt.close(fig)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root',required=True,help='SNR on tissue section data root')
    parser.add_argument('--out',required=True,help='Output directory')
    args=parser.parse_args()
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise SystemExit('Install: python -m pip install openpyxl') from e
    root=Path(args.root).expanduser().resolve()
    out=Path(args.out).expanduser().resolve()
    out.mkdir(parents=True,exist_ok=True)
    rows=[]
    for src in SOURCES:
        path=root/src['relative_path']
        if not path.exists():
            raise SystemExit(f'Excel workbook not found: {path}')
        wb=load_workbook(path,read_only=True,data_only=True)
        try:
            extracted=extract_rows(wb.worksheets[0].iter_rows(values_only=True),src)
        finally:
            wb.close()
        if not extracted:raise SystemExit(f'No numeric measurements found: {path}')
        ids=[r['measurement_id'] for r in extracted]
        if len(set(ids))!=len(ids):
            raise SystemExit(f'Duplicate measurement IDs in {path}')
        rows.extend(extracted)
        print(f"Read {len(extracted)} measurements: {src['condition']} / {src['mode']}")
    summary=summarize(rows)
    write_csv(out/'measurement_audit.csv', rows, list(rows[0]))
    write_csv(out/'group_summary.csv', summary, list(summary[0]))
    plot(rows,summary,out)
    lines=['Chapter 2 historical SNR benchmark (DESCRIPTIVE, NOT CAUSAL)',
           'Reported SNR is historically recorded Signal/Noise; NOT modern CNR based on background SD.', '']
    for c in dict.fromkeys(s['condition'] for s in summary):
        cy=next(s for s in summary if s['condition']==c and s['modality']=='Cy3')
        eu=next(s for s in summary if s['condition']==c and s['modality']=='Eu')
        lines += [f"{c}: Cy3 {cy['reported_snr_mean']:.3f} ± {cy['reported_snr_sd']:.3f} (n={cy['n_measurements']}); "
                  f"Eu {eu['reported_snr_mean']:.3f} ± {eu['reported_snr_sd']:.3f} (n={eu['n_measurements']}); "
                  f"ratio of group means {eu['reported_snr_mean']/cy['reported_snr_mean']:.3f}x"]
    lines += ['', 'AUDIT: rounding consistency checks against Signal / Noise:']
    for s in summary:
        lines.append(f"  {s['condition']} / {s['modality']}: {s['rounding_inconsistencies']} rows inconsistent with usual displayed rounding; unique noise values {s['noise_unique_values']}")
    lines += ['', 'LIMITATIONS:',
              '- The ten measurements are not verified as independent biological samples. Do NOT claim n=10 biological replicates or calculate significance from them.',
              '- Excel measurement provenance, region selection, matched tissue/antibody, and FFEP/FS terminology require confirmation.',
              '- Eu BT474 noise may be fixed at 30; verify original acquisition and quantification method.',
              '- Absolute Eu and Cy3 signal intensities are on different scales; compare only reported ratios IF SNR definitions/protocols match.',
              '- Higher measured ratio is descriptive; it does not isolate the causal effect of time gating without appropriate gated/ungated controls.',
              '- Retain Excel reported SNR and recomputed Signal/Noise separately: some values do not agree after normal decimal rounding.']
    (out/'interpretation.txt').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print('\n'.join(lines));print(f'\nSaved results to {out}')

if __name__=='__main__':main()
